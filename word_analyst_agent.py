"""
Word Document Analyst Agent
Multi-agent architecture for .docx analysis, classification, RAG, contract/resume/research workflows,
table analysis, OCR, editing, generation, and supervision.
"""

from __future__ import annotations

import base64
import json
import os
import re
import time
import zipfile
import difflib
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from dotenv import load_dotenv
from langchain_openai import ChatOpenAI
from langchain_experimental.agents import create_pandas_dataframe_agent
from langchain_classic.memory import ConversationBufferMemory

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover
    DocxDocument = None

try:
    import pytesseract
except ImportError:  # pragma: no cover
    pytesseract = None

try:
    from PIL import Image
except ImportError:  # pragma: no cover
    Image = None

load_dotenv()

DEFAULT_MODEL = os.environ.get("DEEPSEEK_MODEL", os.environ.get("LLM_MODEL", "deepseek-chat"))
DEFAULT_API_KEY = os.environ.get("DEEPSEEK_API_KEY")
DEFAULT_API_BASE = os.environ.get("DEEPSEEK_API_BASE", "https://api.deepseek.com/v1")

EMAIL_RE = re.compile(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[ -.]?)?(?:\(?\d{2,4}\)?[ -.]?)?\d{3,4}[ -.]?\d{3,4}")
MONEY_RE = re.compile(r"\$?\d{1,3}(?:[,.]\d{3})*(?:\.\d+)?\s?(?:USD|EUR|GBP|ETB)?")
URL_RE = re.compile(r"https?://[\w./?=&%-]+")
DATE_RE = re.compile(r"\b(?:\d{1,2}[-/]\d{1,2}[-/]\d{2,4}|\d{4}-\d{2}-\d{2})\b")
ADDRESS_RE = re.compile(
    r"\b\d+\s+[A-Za-z0-9.,'\-\s]+(?:Street|St|Avenue|Ave|Road|Rd|Boulevard|Blvd|Lane|Ln|Drive|Dr|Court|Ct)\b",
    re.IGNORECASE,
)
NAME_RE = re.compile(r"\b[A-Z][a-z]+\s+[A-Z][a-z]+\b")


def build_llm() -> ChatOpenAI:
    kwargs: Dict[str, Any] = {"model": DEFAULT_MODEL, "temperature": 0}
    if DEFAULT_API_KEY:
        kwargs["api_key"] = DEFAULT_API_KEY
    if DEFAULT_API_BASE:
        kwargs["base_url"] = DEFAULT_API_BASE
    return ChatOpenAI(**kwargs)


def retry(max_attempts: int = 3, backoff: float = 0.8):
    def decorator(fn):
        def wrapper(*args: Any, **kwargs: Any) -> Any:
            attempt = 0
            while True:
                try:
                    return fn(*args, **kwargs)
                except Exception:
                    attempt += 1
                    if attempt >= max_attempts:
                        raise
                    time.sleep(backoff * (2 ** (attempt - 1)))

        return wrapper

    return decorator


def llm_call(llm: ChatOpenAI, prompt: str, memory: Optional[ConversationBufferMemory] = None) -> str:
    if memory is not None:
        prompt = "Previous conversation:\n" + getattr(memory, "buffer", "") + "\n" + prompt
    if hasattr(llm, "predict"):
        return llm.predict(prompt)
    if hasattr(llm, "generate"):
        output = llm.generate([prompt])
        return getattr(output, "generations", [[None]])[0][0].text
    return llm(prompt)


def extract_entities(text: str) -> Dict[str, List[str]]:
    return {
        "emails": list(set(EMAIL_RE.findall(text))),
        "phones": list(set(PHONE_RE.findall(text))),
        "money": list(set(MONEY_RE.findall(text))),
        "urls": list(set(URL_RE.findall(text))),
        "dates": list(set(DATE_RE.findall(text))),
        "addresses": list(set(ADDRESS_RE.findall(text))),
    }


def redact_text(text: str, targets: List[str]) -> str:
    redacted = text
    if "emails" in targets:
        redacted = EMAIL_RE.sub("[REDACTED EMAIL]", redacted)
    if "phones" in targets:
        redacted = PHONE_RE.sub("[REDACTED PHONE]", redacted)
    if "addresses" in targets:
        redacted = ADDRESS_RE.sub("[REDACTED ADDRESS]", redacted)
    if "urls" in targets:
        redacted = URL_RE.sub("[REDACTED URL]", redacted)
    if "names" in targets:
        redacted = NAME_RE.sub("[REDACTED NAME]", redacted)
    return redacted


def text_similarity(text_a: str, text_b: str) -> float:
    return round(difflib.SequenceMatcher(None, text_a, text_b).ratio() * 100, 2)


def should_generate_csv(prompt: str) -> bool:
    normalized = prompt.lower()
    return "convert to csv" in normalized or "to csv" in normalized


def should_generate_excel(prompt: str) -> bool:
    normalized = prompt.lower()
    return "convert to excel" in normalized or "to excel" in normalized


def create_excel_base64(df: pd.DataFrame) -> Optional[str]:
    try:
        buffer = BytesIO()
        with pd.ExcelWriter(buffer, engine="xlsxwriter") as writer:
            df.to_excel(writer, index=False)
        buffer.seek(0)
        return base64.b64encode(buffer.read()).decode("utf-8")
    except Exception:
        return None


def parse_tabular_input(csv_content: Optional[str], table_json: Optional[List[Dict[str, Any]]]) -> pd.DataFrame:
    if csv_content:
        try:
            return pd.read_csv(StringIO(csv_content))
        except Exception as exc:
            raise ValueError(f"Failed to parse CSV content: {exc}") from exc

    if table_json is not None:
        try:
            return pd.DataFrame(table_json)
        except Exception as exc:
            raise ValueError(f"Failed to parse table_json: {exc}") from exc

    return pd.DataFrame()


def build_docx_tables(table_frames: List[pd.DataFrame]) -> pd.DataFrame:
    """Combine multiple table DataFrames from a Word document into a single DataFrame."""
    if not table_frames:
        return pd.DataFrame()
    if len(table_frames) == 1:
        return table_frames[0]
    try:
        return pd.concat(table_frames, axis=0, ignore_index=True)
    except Exception:
        return table_frames[0] if table_frames else pd.DataFrame()


def build_doc_summary(df: pd.DataFrame) -> Dict[str, Any]:
    summary: Dict[str, Any] = {
        "total_documents": len(df),
        "columns": df.columns.tolist(),
        "rows": len(df),
    }
    if "content" in df.columns:
        summary["document_names"] = (
            df["document_name"].astype(str).tolist()
            if "document_name" in df.columns
            else [f"doc_{i+1}" for i in range(len(df))]
        )
        summary["word_counts"] = df["content"].astype(str).str.split().apply(len).tolist()
        summary["avg_words_per_doc"] = float(df["content"].astype(str).str.split().apply(len).mean())
        if "sections" in df.columns:
            summary["sections"] = df["sections"].astype(str).tolist()
    summary["languages_detected"] = detect_languages_from_texts(df["content"].astype(str).tolist()) if "content" in df.columns else []
    return summary


def detect_languages_from_texts(texts: List[str]) -> List[str]:
    languages = set()
    for text in texts:
        lang = detect_language(text)
        if lang:
            languages.add(lang)
    return sorted(languages)


def detect_language(text: str) -> str:
    if re.search(r"[\u1200-\u137F]", text):
        return "Amharic"
    if re.search(r"[\u0400-\u04FF]", text):
        return "Cyrillic"
    if re.search(r"[\u0600-\u06FF]", text):
        return "Arabic"
    if re.search(r"[\u4E00-\u9FFF]", text):
        return "Chinese"
    if re.search(r"[\u0900-\u097F]", text):
        return "Hindi"
    if re.search(r"[\u0400-\u052F]", text):
        return "Russian"
    if re.search(r"[\u0000-\u007F]", text) and not re.search(r"[\u1200-\u137F\u0400-\u04FF\u0600-\u06FF\u4E00-\u9FFF\u0900-\u097F]", text):
        return "English"
    return "Unknown"


def extract_section_diffs(df: pd.DataFrame) -> Dict[str, List[str]]:
    diffs = {"added_sections": [], "removed_sections": [], "changed_clauses": [], "risk_changes": []}
    if "document_name" not in df.columns or "sections" not in df.columns or len(df) < 2:
        return diffs

    docs = df.groupby("document_name")["sections"].agg(lambda x: ", ".join(x.astype(str))).to_dict()
    if len(docs) < 2:
        return diffs

    names = list(docs.keys())[:2]
    sections_a = set(re.split(r"[;\n]+", docs[names[0]]))
    sections_b = set(re.split(r"[;\n]+", docs[names[1]]))
    diffs["added_sections"] = sorted([s.strip() for s in sections_b - sections_a if s.strip()])
    diffs["removed_sections"] = sorted([s.strip() for s in sections_a - sections_b if s.strip()])
    return diffs


class OCRAgent:
    def extract_image_ocr(self, image_bytes: bytes, filename: str) -> str:
        if pytesseract is None or Image is None:
            return ""
        try:
            with Image.open(BytesIO(image_bytes)) as image:
                return pytesseract.image_to_string(image, lang="eng").strip()
        except Exception:
            return ""

    def extract_docx_images(self, file_path: str) -> List[Dict[str, Any]]:
        images: List[Dict[str, Any]] = []
        path = Path(file_path)
        if not path.exists() or path.suffix.lower() != ".docx":
            return images

        if pytesseract is None or Image is None:
            return images

        try:
            with zipfile.ZipFile(str(path), "r") as archive:
                for member in archive.namelist():
                    if member.startswith("word/media/") and not member.endswith("/"):
                        image_bytes = archive.read(member)
                        ocr_text = self.extract_image_ocr(image_bytes, member)
                        images.append({"name": member, "bytes_length": len(image_bytes), "ocr_text": ocr_text})
        except Exception:
            return images

        return images


class DocumentReaderAgent:
    def __init__(self) -> None:
        self.ocr_agent = OCRAgent()

    def extract_docx_content(self, file_path: str) -> Dict[str, Any]:
        if DocxDocument is None:
            raise RuntimeError("python-docx is required to extract .docx files")

        path = Path(file_path)
        if not path.exists() or path.suffix.lower() != ".docx":
            raise ValueError("extract_docx_content requires a valid .docx file path")

        document = DocxDocument(str(path))
        paragraphs: List[str] = []
        headings: List[Dict[str, Any]] = []
        tables: List[Dict[str, Any]] = []
        headers: List[str] = []
        footers: List[str] = []
        styles: List[Dict[str, str]] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style is not None else "Normal"
            if style_name.lower().startswith("heading"):
                level_match = re.search(r"\d+", style_name)
                level = int(level_match.group(0)) if level_match else 1
                headings.append({"text": text, "level": level, "style": style_name})
            else:
                paragraphs.append(text)

        for table in document.tables:
            tables.append(self._table_to_struct(table))

        for section in document.sections:
            if section.header is not None:
                header_text = "\n".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
                if header_text:
                    headers.append(header_text)
            if section.footer is not None:
                footer_text = "\n".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
                if footer_text:
                    footers.append(footer_text)

        for style in document.styles:
            styles.append({"name": style.name, "type": str(style.type)})

        metadata = self._extract_metadata(document)
        images = self.ocr_agent.extract_docx_images(file_path)
        ocr_texts = [image["ocr_text"] for image in images if image.get("ocr_text")]

        text_fragments: List[str] = paragraphs.copy()
        if ocr_texts:
            text_fragments.append("\n\n".join(ocr_texts))

        return {
            "text": "\n\n".join(text_fragments),
            "paragraphs": paragraphs,
            "headings": headings,
            "tables": tables,
            "headers": headers,
            "footers": footers,
            "styles": styles,
            "images": images,
            "ocr_text": "\n\n".join(ocr_texts),
            "ocr_results": images,
            "metadata": metadata,
        }

    def _extract_metadata(self, document: Any) -> Dict[str, Any]:
        metadata: Dict[str, Any] = {}
        try:
            from datetime import datetime
            core_props = document.core_properties
            for field in [
                "author",
                "category",
                "comments",
                "content_status",
                "created",
                "identifier",
                "keywords",
                "language",
                "last_modified_by",
                "last_printed",
                "modified",
                "revision",
                "subject",
                "title",
                "version",
            ]:
                value = getattr(core_props, field, None)
                if value is not None:
                    # Convert datetime objects to ISO format strings for JSON serialization
                    if isinstance(value, datetime):
                        metadata[field] = value.isoformat()
                    else:
                        metadata[field] = value
        except Exception:
            metadata["error"] = "Unable to extract core document properties"
        return metadata

    def _table_to_struct(self, table: Any) -> Dict[str, Any]:
        rows: List[List[str]] = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        headers = rows[0] if rows else []
        body = rows[1:] if len(rows) > 1 else []
        return {"headers": headers, "rows": body, "row_count": len(body)}


class DocumentClassifierAgent:
    def classify(self, full_text: str, metadata: Dict[str, Any]) -> Dict[str, Any]:
        normalized = full_text.lower()
        if any(keyword in normalized for keyword in ["contract", "agreement", "clause", "party", "warranty"]):
            return {"type": "contract", "confidence": 0.95}
        if any(keyword in normalized for keyword in ["resume", "cv", "curriculum vitae", "education", "experience"]):
            return {"type": "resume", "confidence": 0.95}
        if any(keyword in normalized for keyword in ["research", "methodology", "results", "findings", "abstract"]):
            return {"type": "research", "confidence": 0.95}
        if any(keyword in normalized for keyword in ["invoice", "billing", "payment", "amount", "due"]):
            return {"type": "financial", "confidence": 0.85}
        if any(keyword in normalized for keyword in ["proposal", "project", "plan", "summary"]):
            return {"type": "business", "confidence": 0.8}
        return {"type": "general", "confidence": 0.6}


class RAGAgent:
    def __init__(self, llm: ChatOpenAI, document_text: str, docx_data: Optional[Dict[str, Any]] = None) -> None:
        self.llm = llm
        self.document_text = document_text or ""
        self.docx_data = docx_data or {}
        self.chunks = self._chunk_text(self.document_text)

    def _chunk_text(self, text: str, chunk_size: int = 1000, overlap: int = 200) -> List[str]:
        if not text:
            return []
        paragraphs = [p.strip() for p in re.split(r"\n{2,}|\r\n{2,}", text) if p.strip()]
        chunks: List[str] = []
        current = ""
        for paragraph in paragraphs:
            if len(current) + len(paragraph) + 2 > chunk_size and current:
                chunks.append(current.strip())
                current = paragraph
            else:
                current = f"{current}\n\n{paragraph}".strip()
            if len(current) >= chunk_size:
                chunks.append(current.strip())
                current = ""
        if current:
            chunks.append(current.strip())
        return chunks

    def _score(self, query: str, chunk: str) -> float:
        return difflib.SequenceMatcher(None, query.lower(), chunk.lower()).ratio()

    def retrieve_top_chunks(self, query: str, top_n: int = 3) -> List[str]:
        scored = [(self._score(query, chunk), chunk) for chunk in self.chunks]
        scored.sort(key=lambda item: item[0], reverse=True)
        return [chunk for _, chunk in scored[:top_n] if chunk]

    def answer_query(self, query: str, memory: Optional[ConversationBufferMemory] = None) -> Dict[str, Any]:
        if not query:
            return {"tool": "semantic_search", "search_result": "No query provided."}
        context = "\n\n".join(self.retrieve_top_chunks(query, top_n=3))
        prompt = (
            "Use the following document context to answer the user's query as precisely as possible. "
            "If the answer is not contained in the context, say so clearly.\n\n"
            f"Query:\n{query}\n\nContext:\n{context}"
        )
        answer = llm_call(self.llm, prompt, memory=memory)
        return {"tool": "semantic_search", "search_result": answer, "context_used": context}


class ContractAnalysisAgent:
    def analyze(self, full_text: str, llm: ChatOpenAI) -> Dict[str, Any]:
        prompt = (
            "Analyze the contract text and identify key obligations, risks, clauses, and potential issues. "
            f"\n\nContract text:\n{full_text[:4000]}"
        )
        return {"tool": "contract_analysis", "analysis": llm_call(llm, prompt)}


class ResumeAnalysisAgent:
    def analyze(self, full_text: str, llm: ChatOpenAI) -> Dict[str, Any]:
        prompt = (
            "Review this resume content and summarize the candidate's strengths, experience, and fit. "
            "Include recommended role focus areas."
            f"\n\nResume text:\n{full_text[:4000]}"
        )
        return {"tool": "resume_analysis", "analysis": llm_call(llm, prompt)}


class ResearchAnalysisAgent:
    def analyze(self, full_text: str, llm: ChatOpenAI) -> Dict[str, Any]:
        prompt = (
            "Review the research document and summarize the key findings, methodology, and conclusions. "
            "Highlight any gaps or opportunities."
            f"\n\nResearch text:\n{full_text[:4000]}"
        )
        return {"tool": "research_analysis", "analysis": llm_call(llm, prompt)}


class TableAnalysisAgent:
    def analyze(self, df: pd.DataFrame, prompt: str = "", table_agent: Any = None) -> Dict[str, Any]:
        if df.empty:
            return {"tool": "table_analysis", "error": "No table data available."}
        if table_agent is not None:
            try:
                instruction = prompt or "Summarize the table data."
                result = table_agent.invoke(instruction) if hasattr(table_agent, "invoke") else table_agent.run(instruction)
                return {"tool": "table_analysis", "table_agent_result": result}
            except Exception as exc:
                return {"tool": "table_analysis", "error": f"Table agent failed: {exc}"}
        return {"tool": "table_analysis", "summary": build_doc_summary(df)}


class DocumentEditingAgent:
    def edit(self, full_text: str, prompt: str, llm: ChatOpenAI, file_path: Optional[str] = None) -> Dict[str, Any]:
        if file_path and Path(file_path).suffix.lower() == ".docx" and DocxDocument is not None:
            return self._edit_docx(file_path, prompt)

        answer = llm_call(
            llm,
            "Edit the following document according to the user request, preserving tone and structure. "
            f"\n\nRequest:\n{prompt}\n\nDocument:\n{full_text[:4000]}",
        )
        return {"tool": "edit_document", "edited_document": answer}

    def _edit_docx(self, file_path: str, prompt: str) -> Dict[str, Any]:
        original_path = Path(file_path)
        if not original_path.exists() or original_path.suffix.lower() != ".docx":
            return {"tool": "edit_document", "error": "A valid .docx file path is required for editing."}

        reader = DocumentReaderAgent()
        docx_data = reader.extract_docx_content(file_path)
        original_text = docx_data.get("text", "")
        llm = build_llm()
        editing_prompt = (
            "Rewrite the following document to satisfy the user's request while preserving the original meaning. "
            "Return only the revised text with paragraph breaks preserved.\n\n"
            f"Request:\n{prompt}\n\nDocument text:\n{original_text}"
        )
        edited_text = llm_call(llm, editing_prompt)

        edited_doc_path = original_path.with_name(f"{original_path.stem}_edited.docx")
        edited_document = DocxDocument()
        for paragraph_text in edited_text.splitlines():
            edited_document.add_paragraph(paragraph_text)
        edited_document.save(str(edited_doc_path))

        return {
            "tool": "edit_document",
            "edited_docx_path": str(edited_doc_path),
            "edited_text": edited_text,
            "original_docx_path": str(original_path),
        }


class DocumentGenerationAgent:
    def generate(self, prompt: str, llm: ChatOpenAI) -> Dict[str, Any]:
        answer = llm_call(llm, f"Generate a document based on this request:\n{prompt}")
        return {"tool": "generate_document", "generated_document": answer}


class SupervisorAgent:
    def __init__(self) -> None:
        self.llm = build_llm()
        self.reader = DocumentReaderAgent()
        self.classifier = DocumentClassifierAgent()
        self.contract_agent = ContractAnalysisAgent()
        self.resume_agent = ResumeAnalysisAgent()
        self.research_agent = ResearchAnalysisAgent()
        self.table_agent = TableAnalysisAgent()
        self.editor = DocumentEditingAgent()
        self.generator = DocumentGenerationAgent()
        self.rag: Optional[RAGAgent] = None

    def detect_intent(self, prompt: str) -> str:
        normalized = prompt.lower()
        if any(keyword in normalized for keyword in ["contract", "agreement", "terms", "clause", "warranty", "party"]):
            return "contract_analysis"
        if any(keyword in normalized for keyword in ["resume", "cv", "curriculum vitae", "experience", "skills", "education"]):
            return "resume_analysis"
        if any(keyword in normalized for keyword in ["research", "abstract", "methodology", "results", "findings", "literature review"]):
            return "research_analysis"
        if any(keyword in normalized for keyword in ["table", "spreadsheet", "csv", "excel", "data frame", "calculate", "analysis"]):
            return "table_analysis"
        if any(keyword in normalized for keyword in ["edit document", "revise", "rewrite", "improve", "modify"]):
            return "edit_document"
        if any(keyword in normalized for keyword in ["generate document", "create document", "draft", "write a document", "compose"]):
            return "generate_document"
        if any(keyword in normalized for keyword in ["metadata", "properties", "document info", "document metadata", "title", "author"]):
            return "metadata_analysis"
        if any(keyword in normalized for keyword in ["search", "find", "lookup", "locate", "where is", "what section"]):
            return "semantic_search"
        if any(keyword in normalized for keyword in ["redact", "sanitized", "sanitize", "mask", "obfuscate"]):
            return "redact"
        return "summarize"

    def _build_table_agent(self, df: pd.DataFrame) -> Any:
        if df.empty:
            return None
        return create_pandas_dataframe_agent(
            llm=self.llm,
            df=df,
            verbose=False,
            agent_type="tool-calling",
            allow_dangerous_code=True,
            number_of_head_rows=5,
        )

    def route_tool(
        self,
        intent: str,
        full_text: str,
        prompt: str,
        df: pd.DataFrame,
        table_agent: Any,
        metadata: Dict[str, Any],
        classification: Dict[str, Any],
        memory: ConversationBufferMemory,
        file_path: Optional[str],
    ) -> Dict[str, Any]:
        if intent == "contract_analysis":
            return self.contract_agent.analyze(full_text, self.llm)
        if intent == "resume_analysis":
            return self.resume_agent.analyze(full_text, self.llm)
        if intent == "research_analysis":
            return self.research_agent.analyze(full_text, self.llm)
        if intent == "table_analysis":
            return self.table_agent.analyze(df, prompt=prompt, table_agent=table_agent)
        if intent == "edit_document":
            return self.editor.edit(full_text, prompt, self.llm, file_path=file_path)
        if intent == "generate_document":
            return self.generator.generate(prompt, self.llm)
        if intent == "semantic_search":
            if self.rag is None:
                self.rag = RAGAgent(self.llm, full_text)
            return self.rag.answer_query(prompt, memory=memory)
        if intent == "metadata_analysis":
            return {"tool": "metadata_analysis", "metadata": metadata, "classification": classification}
        if intent == "redact":
            return {"tool": "redact", "redacted_text": redact_text(full_text, ["emails", "phones", "addresses", "urls", "names"])}
        return {"tool": "summarize", "summary": self._summarize_document(full_text, prompt)}

    def _summarize_document(self, full_text: str, prompt: str) -> str:
        if not full_text:
            return "No document text available to summarize."
        summary_prompt = (
            "Summarize the document text in a clear way, highlighting the main points, structure, and key findings. "
            f"Use the user request: {prompt}\n\nText:\n{full_text[:4000]}"
        )
        return llm_call(self.llm, summary_prompt)

    def _prepare_document(self, csv_content: Optional[str], table_json: Optional[List[Dict[str, Any]]], file_path: Optional[str]) -> Tuple[pd.DataFrame, Optional[Dict[str, Any]], Dict[str, Any]]:
        df = parse_tabular_input(csv_content, table_json)
        docx_data = None
        if file_path:
            docx_data = self.reader.extract_docx_content(file_path)
            table_frames: List[pd.DataFrame] = []
            for table in docx_data.get("tables", []):
                if table.get("rows"):
                    headers = table.get("headers") or []
                    rows = table.get("rows")
                    if headers and len(headers) == len(rows[0]):
                        table_frames.append(pd.DataFrame(rows, columns=headers))
                    else:
                        normalized_rows = [dict(zip(headers, row)) if headers else {str(idx): value for idx, value in enumerate(row)} for row in rows]
                        table_frames.append(pd.DataFrame(normalized_rows))
            if df.empty and table_frames:
                df = build_docx_tables(table_frames)

        metadata: Dict[str, Any] = {"total_documents": len(df), "columns": df.columns.tolist() if not df.empty else []}
        if docx_data is not None:
            metadata["docx"] = {k: v for k, v in docx_data.items() if k in {"headers", "footers", "styles", "metadata"}}

        return df, docx_data, metadata

    def _build_overview_answer(
        self,
        full_text: str,
        metadata: Dict[str, Any],
        entities: Dict[str, List[str]],
        prompt: str,
        memory: ConversationBufferMemory,
    ) -> str:
        prompt_parts = [
            "You are a document analysis assistant.",
            "Use the provided data, document content, and tables to answer the user's task.",
            f"Task: {prompt}",
            f"Metadata: {json.dumps(metadata, ensure_ascii=False)}",
        ]
        if full_text:
            prompt_parts.append(f"Document text sample:\n{full_text[:4000]}")
        prompt_parts.append(f"Detected entities: {json.dumps(entities, ensure_ascii=False)}")
        prompt_parts.append(
            "If the task requires CSV or table calculations, respond with TABLE_AGENT_AVAILABLE: followed by the command you want the agent to run."
        )
        return llm_call(self.llm, "\n\n".join(prompt_parts), memory=memory)

    def run(self, csv_content: Optional[str] = None, table_json: Optional[List[Dict[str, Any]]] = None, prompt: str = "", file_path: Optional[str] = None) -> Dict[str, Any]:
        try:
            df, docx_data, metadata = self._prepare_document(csv_content, table_json, file_path)
            if docx_data is not None:
                full_text = docx_data.get("text", "")
            elif not df.empty and "content" in df.columns:
                full_text = "\n\n".join(df["content"].astype(str).tolist())
            else:
                full_text = ""

            entities = extract_entities(full_text)
            table_agent = self._build_table_agent(df)
            memory = ConversationBufferMemory()
            overview_answer = self._build_overview_answer(full_text, metadata, entities, prompt, memory)

            table_agent_output = None
            if table_agent is not None and "TABLE_AGENT_AVAILABLE" in overview_answer:
                instruction = "".join(overview_answer.split("TABLE_AGENT_AVAILABLE", 1)[1:]).strip().lstrip(":")
                instruction = instruction or prompt
                try:
                    table_agent_output = (
                        table_agent.invoke(instruction)
                        if hasattr(table_agent, "invoke")
                        else table_agent.run(instruction)
                    )
                except Exception as exc:
                    table_agent_output = f"Table agent failed: {exc}"

            classification = self.classifier.classify(full_text, metadata)
            metadata["classification"] = classification
            if full_text and "languages_detected" not in metadata:
                metadata["languages_detected"] = detect_languages_from_texts([full_text])
            intent = self.detect_intent(prompt)
            self.rag = RAGAgent(self.llm, full_text, docx_data)
            tool_result = self.route_tool(intent, full_text, prompt, df, table_agent, metadata, classification, memory, file_path)

            redacted_text = None
            if "redact" in prompt.lower() or "sanitized" in prompt.lower():
                redacted_text = redact_text(full_text, ["emails", "phones", "addresses", "urls", "names"])

            similarity_score = None
            section_diffs: Dict[str, List[str]] = {"added_sections": [], "removed_sections": [], "changed_clauses": [], "risk_changes": []}
            if "compare" in prompt.lower() and len(df) >= 2:
                if "document_name" in df.columns:
                    doc_texts = df.groupby("document_name")["content"].apply(lambda x: "\n\n".join(x.astype(str))).to_dict()
                    names = list(doc_texts.keys())
                    if len(names) >= 2:
                        similarity_score = text_similarity(doc_texts[names[0]], doc_texts[names[1]])
                        section_diffs = extract_section_diffs(df)

            csv_export = None
            excel_export_base64 = None
            if should_generate_csv(prompt) and not df.empty:
                csv_export = df.to_csv(index=False)
            if should_generate_excel(prompt) and not df.empty:
                excel_export_base64 = create_excel_base64(df)

            result: Dict[str, Any] = {
                "answer": overview_answer,
                "intent": intent,
                "tool_routing": tool_result,
                "entities": entities,
                "metadata": metadata,
                "tables_present": not df.empty,
                "table_agent_output": table_agent_output,
                "classification": classification,
                "similarity_score": similarity_score,
                "section_diffs": section_diffs,
                "redacted_text": redacted_text,
                "csv_export": csv_export,
                "excel_export_base64": excel_export_base64,
                "memory_buffer": getattr(memory, "buffer", None),
                "document": docx_data,
            }

            if not df.empty:
                result["summary_stats"] = {
                    "total_words": int(df["content"].astype(str).str.split().apply(len).sum()) if "content" in df.columns else 0,
                    "avg_words_per_doc": float(df["content"].astype(str).str.split().apply(len).mean()) if "content" in df.columns else 0,
                    "top_sections": df["sections"].value_counts().head(5).to_dict() if "sections" in df.columns else {},
                }

            return result
        except Exception as exc:
            return {"error": str(exc), "agent": "word-analyst"}


@retry(max_attempts=3)
def run_word_analyst(
    csv_content: Optional[str] = None,
    table_json: Optional[List[Dict[str, Any]]] = None,
    prompt: str = "",
    file_path: Optional[str] = None,
) -> Dict[str, Any]:
    supervisor = SupervisorAgent()
    return supervisor.run(csv_content=csv_content, table_json=table_json, prompt=prompt, file_path=file_path)
